"""Inserta las respuestas de las Tasks 1 y 2 en el docx del laboratorio 10 (parte de la copia de la plantilla).
Uso: python tools/fill_docx.py <plantilla.docx> <destino.docx> <raíz del repo>
Texto: **negrita**, $latex$ en línea; c.eq(latex) para ecuaciones centradas."""
import json, math, sys
import docx
from docx_helpers import Cursor, style_run

SRC, DST, ROOT = sys.argv[1], sys.argv[2], sys.argv[3]
R = json.load(open(f"{ROOT}/results/results_t1_t2.json", encoding="utf8"))
doc = docx.Document(SRC)
P = list(doc.paragraphs)                       # anclas por índice (antes de insertar nada)
C = lambda i: Cursor(doc, P[i], f"{ROOT}/figures")
f2 = lambda x: f"{x:.2f}"; f3 = lambda x: f"{x:.3f}"; f4 = lambda x: f"{x:.4f}"
sci = lambda x: f"{x:.2e}"

def code(c, name, n=0, w=6.6):
    c.image(f"code/{name}_{n}.png", w)

# ============================================================ encabezado
style_run(P[2].add_run(" https://github.com/Branuvg/labs_deep/tree/lab10"))
c = C(2)
c.para("Semilla aleatoria **SEED = 42**, fijada en random, numpy y torch (CPU y CUDA); el fondo de SHAP, el ruido del Task 2.3 y los muestreos de KernelSHAP "
       "parten de esa misma semilla. Código: lab10.ipynb. Entorno: Python 3.12 (venv), PyTorch 2.11 + CUDA 12.8 en una GPU NVIDIA RTX 5060 Laptop, "
       "scikit-learn 1.9.1 y shap 0.53.0.", indent=False)

# ============================================================ Task 1.1
t = R["t1_1"]
c = C(14)
c.para("**Modelo.** ResNet-18 de torchvision con los pesos IMAGENET1K_V1, en modo eval. Preprocesado: Resize(256) del lado corto, CenterCrop(224) y normalización por canal "
       "con media (0.485, 0.456, 0.406) y desviación estándar (0.229, 0.224, 0.225).")
c.para("**Imágenes y fuentes** (todas de Wikimedia Commons; se recortó previamente un cuadrado alrededor del objeto para que el recorte central de 224×224 no lo cortara):")
c.para("(a) “Golden retriever portrait.jpg”, CC BY-SA 3.0, commons.wikimedia.org/wiki/File:Golden_retriever_portrait.jpg. Es un primer plano de un perro, sin recorte adicional. "
       "El título de Commons dice golden retriever, pero el modelo lo ve como Great Pyrenees; el golden retriever queda tercero.", indent=True)
c.para("(b) “Animals-gfce1d0472 1280.jpg”, CC0, commons.wikimedia.org/wiki/File:Animals-gfce1d0472_1280.jpg. Muestra un gato (a la izquierda) y un cachorro de perro (a la derecha). "
       "Se recortó el cuadrado de 853×853 con x ∈ [100, 953] para que ambos animales entraran en el recorte central.", indent=True)
c.para("(c) “Abandoned sofa, Picket Post viewpoint, New Forest” (geograph.org.uk, imagen 692118), CC BY-SA 2.0, "
       "commons.wikimedia.org/wiki/File:Abandoned_sofa,_Picket_Post_viewpoint,_New_Forest_-_geograph.org.uk_-_692118.jpg. Es un sofá blanco abandonado sobre un campo de pasto, "
       "un objeto de interior en un contexto de exterior. Se recortó el cuadrado de 427×427 con x ∈ [213, 640].", indent=True)
for k, nm in [("a", "(a) objeto centrado"), ("b", "(b) gato y perro"), ("c", "(c) sofá en el campo")]:
    c.table(["#", "Clase ImageNet", "Índice", "Probabilidad"],
            [[i + 1, r["name"], r["idx"], f4(r["p"])] for i, r in enumerate(t[k])], f"Tabla 1{k}. Cinco clases de mayor probabilidad, imagen {nm}.")
x0, y0, x1, y1 = t["box_c"]
c.para(f"**Recuadro de la imagen (c)**, medido sobre la imagen de 224×224: esquina superior izquierda $(x_0,y_0)=({x0},{y0})$ y esquina inferior derecha $(x_1,y_1)=({x1},{y1})$ "
       f"(ancho {x1-x0}, alto {y1-y0}); cubre el {100*t['box_c_area_frac']:.1f}% de la imagen.")
c.image("t1_1_imagenes.png", 6.4, "Figura 1. Las tres imágenes de 224×224 tras el preprocesado (el recuadro verde es el de la imagen c).")
c.para(f"En (a) el modelo asigna {100*t['a'][0]['p']:.0f}% a Great Pyrenees y el resto a razas de aspecto parecido (kuvasz, golden retriever, Samoyed, malamute): reconoce con claridad un perro grande de pelaje claro. En (b) el perro domina: Labrador retriever {f3(t['b'][0]['p'])}, "
       f"mientras que el gato solo aparece en 3.er lugar como tiger cat ({f3(t['b'][2]['p'])}). En (c) el modelo no reconoce un sofá: predice bathtub ({f3(t['c'][0]['p'])}) y tub ({f3(t['c'][1]['p'])}), "
       "una forma blanca y redondeada que confunde con una bañera.")
code(c, "t1_1", 0); code(c, "t1_1", 1)

# ============================================================ Task 1.2
g = R["t1_2"]
c = C(21)
c.para("Se implementó la clase GradCAM del notebook con un hook de avance (register_forward_hook) que guarda $A$ (512×7×7) y un hook de retroceso (register_full_backward_hook) que guarda $\\partial y^c/\\partial A$ "
       "en layer4. La retropropagación sale del **logit** $y^c$ (logits[0, c].backward()), sin softmax ni pérdida. Los pasos son:")
c.eq(r"\alpha_k^c=\frac{1}{Z}\sum_{i,j}\frac{\partial y^c}{\partial A^k_{ij}},\qquad L^c=\mathrm{ReLU}\Big(\sum_k\alpha_k^cA^k\Big)\in\mathbb{R}^{7\times7}")
c.para("Después $L^c$ se amplía a 224×224 con interpolación bilineal (F.interpolate) y se normaliza con $(L-\\min L)/(\\max L-\\min L)$. "
       f"Clases explicadas: (a) {g['a']['name']} (índice {g['a']['cls']}, p = {f3(g['a']['p'])}); (b) {g['b']['name']} (p = {f3(g['b']['p'])}) y, como clase del segundo objeto, "
       f"**{g['b2']['name']}** (índice {g['b2']['cls']}, p = {f4(g['b2']['p'])}), que es la clase de gato más probable (índices 281–285); (c) {g['c']['name']} (p = {f3(g['c']['p'])}).")
c.image("t1_2_gradcam.png", 6.6, "Figura 2. Grad-CAM superpuesto (jet, α = 0.5) para la clase de mayor probabilidad de cada imagen y, en (b), también para la del segundo objeto.")
c.para("**Lectura.** En (a) el mapa se concentra en la cara del perro (ojo, hocico y mejilla) y casi ignora las tablas verdes del fondo. En (b) el mapa de Labrador retriever cubre la cabeza y el cuello del cachorro, "
       "y el de tiger cat se desplaza al cuerpo y la cabeza del gato, en la zona izquierda. En (c) el calor se centra en el sofá, pero forma un halo que se extiende al pasto y la sombra que lo rodean "
       "porque cada celda original cubre 32×32 píxeles y la interpolación bilineal suaviza los bordes.")
code(c, "t1_2")

# ============================================================ Task 1.3
r = R["t1_3"]
c = C(25)
c.para("Con los pesos de fc de la clase $c$ de cada imagen se calcula $M^c_{ij}=\\sum_k w^c_kA^k_{ij}$ (torch.einsum sobre las activaciones guardadas, sin ningún gradiente). "
       "En el notebook se obtuvo para las cuatro clases explicadas (a, b, b2 = segundo objeto, c). Figura 3: mapas 7×7 de CAM (arriba) y de Grad-CAM (abajo), ambos con ReLU y min-max, idénticos a simple vista.")
c.image("t1_3_cam_vs_gradcam.png", 6.4, "Figura 3. CAM calculado con los pesos de fc (fila superior) frente a Grad-CAM (fila inferior).")
code(c, "t1_3")
names = {"a": "(a) " + g["a"]["name"], "b": "(b) " + g["b"]["name"], "b2": "(b) " + g["b2"]["name"], "c": "(c) " + g["c"]["name"]}
c = C(27)
c.para("Como el modelo es pooling global más capa lineal, $y^c=\\frac1Z\\sum_{i,j}M^c_{ij}+b_c$ con $Z=49$:")
c.table(["Imagen / clase", "(1/Z)·ΣM + b_c", "Logit y^c", "Error absoluto"],
        [[names[k], f"{r[k]['score']:.6f}", f"{r[k]['logit']:.6f}", sci(r[k]["err_b"])] for k in ["a", "b", "b2", "c"]], "Tabla 2. Descomposición del puntaje.")
c.para("La igualdad se cumple hasta ≈ $10^{-6}$ en un logit de magnitud 7–14, es decir, hasta el error de redondeo de float32.")
c = C(29)
c.para("Se calculó $L^c$ sin ReLU ni normalización, es decir $\\sum_k\\alpha_k^cA^k$ con $\\alpha$ obtenido por retropropagación, se multiplicó por $Z=49$ y se comparó con $M^c$ en las 49 posiciones:")
c.table(["Imagen / clase", "max |Z·L − M| (49 posiciones)"], [[names[k], sci(r[k]["err_c"])] for k in ["a", "b", "b2", "c"]], "Tabla 3. Relación Grad-CAM–CAM (sin ReLU ni normalización).")
c.para("Los mapas coinciden hasta ≈ $4\\times10^{-6}$, también del orden del redondeo de float32: para este ResNet, Grad-CAM y CAM son el mismo mapa, salvo por un factor $Z$.")
c = C(31)
c.table(["Imagen / clase", "max |error| tras ReLU y min-max"], [[names[k], sci(r[k]["err_d"])] for k in ["a", "b", "b2", "c"]], "Tabla 4. Error entre los mapas normalizados.")
c.para("Tras ReLU y normalización mínimo-máximo el error baja a ≈ $10^{-7}$: la normalización divide por el rango y elimina el factor $Z$, por lo que ambos mapas son indistinguibles.")

# ============================================================ Task 1.4 (a)
a = R["t1_4a"]
c = C(35)
c.para(f"Clase de la imagen (a): {g['a']['name']} (índice {g['a']['cls']}). Se tomaron los cinco canales con mayor $w_k^c$ (los de mayor peso positivo).")
c.table(["Canal k", "w_k^c", "α_k^c (autograd)", "w_k^c / Z", "Error relativo"],
        [[x["canal"], f"{x['w']:.6f}", f"{x['alpha']:.10f}", f"{x['w_sobre_Z']:.10f}", sci(x["err_rel"]) if x["err_rel"] else "0"] for x in a["rows"]],
        "Tabla 5. α obtenido con gradientes frente a w/Z (Z = 49).")
c.para(f"Sobre los 512 canales el error relativo máximo de $\\alpha$ es {sci(a['max_rel_all'])}, y la diferencia máxima $|\\partial y^c/\\partial A^k_{{ij}}-w_k^c/Z|$ en las 49 posiciones de todos los canales es {a['spread']:.1f}.")
c.para("**Paso de la derivación.** Como $y^c=\\sum_kw^c_k\\frac1Z\\sum_{i,j}A^k_{ij}+b_c$ es lineal en $A^k_{ij}$, su derivada es la misma constante en las $Z$ posiciones:")
c.eq(r"\frac{\partial y^c}{\partial A^k_{ij}}=\frac{w^c_k}{Z}\quad\forall\,(i,j)\;\Rightarrow\;\alpha^c_k=\frac{1}{Z}\sum_{i,j}\frac{w_k^c}{Z}=\frac{1}{Z}\cdot Z\cdot\frac{w_k^c}{Z}=\frac{w^c_k}{Z}")
c.para("Se usa al pasar de la definición de $\\alpha_k^c$ (una suma de $Z$ derivadas, una por posición) a una sola constante: la suma sobre $(i,j)$ pasa a ser $Z$ copias del mismo valor y el $1/Z$ del promedio las cancela. "
       "Esto se verifica en los resultados, donde la dispersión espacial del gradiente es exactamente 0.0.")
c.para("**Por qué el error no es exactamente cero.** La igualdad matemática es exacta, pero el cálculo no lo es. Autograd propaga el gradiente de avgpool como $w_k/49$ en cada posición, en float32, y luego "
       "$\\alpha_k$ se obtiene promediando esos 49 valores con sumas redondeadas, mientras que $w_k/Z$ es una sola división. Son dos secuencias de operaciones distintas y cada una acumula un redondeo de ≈ $\\epsilon_{mach}\\approx1.2\\times10^{-7}$ "
       f"(épsilon de float32). Por eso los errores relativos son 0 o ≈ $10^{{-7}}$ (máximo {sci(a['max_rel_all'])} en 512 canales, ≈ 1.5 épsilon) y no 0 exacto.")
code(c, "t1_4a")

# ============================================================ Task 1.4 (b)
b = R["t1_4b"]
c = C(37)
c.table(["Mapa", "Probabilidad", "Centroide (x, y) en 224×224"],
        [[b["cls1"], f3(b["p1"]), f"({b['centroid1'][0]:.0f}, {b['centroid1'][1]:.0f})"], [b["cls2"], f3(b["p2"]), f"({b['centroid2'][0]:.0f}, {b['centroid2'][1]:.0f})"]],
        "Tabla 6. Los dos mapas de la imagen (b).")
c.para(f"Se tomó en cada mapa normalizado el 20% de píxeles de mayor valor (umbral en el percentil 80). **IoU = {b['iou224']:.3f}** sobre la imagen de 224×224 y {b['iou7']:.3f} sobre los mapas 7×7 originales "
       f"(10 de 49 celdas cada uno). La correlación de Pearson entre los dos mapas completos es {b['corr']:.2f}.")
c.image("t1_4b_dos_clases.png", 6.4, "Figura 4. Mapas de las dos clases y sus conjuntos del 20% superior (rojo = clase 1, azul = clase 2).")
c.para(f"**¿Discrimina?** Sí. Los conjuntos del 20% superior son disjuntos (IoU = 0) y el mapa de {b['cls1']} está arriba a la derecha (centroide ≈ ({b['centroid1'][0]:.0f}, {b['centroid1'][1]:.0f}), el cachorro) mientras que el de {b['cls2']} "
       f"está abajo a la izquierda (≈ ({b['centroid2'][0]:.0f}, {b['centroid2'][1]:.0f}), el gato), a ≈ {math.hypot(b['centroid1'][0]-b['centroid2'][0], b['centroid1'][1]-b['centroid2'][1]):.0f} px de distancia. La correlación negativa confirma que son mapas complementarios. "
       f"Cabe una salvedad: la clase tiger cat tiene p = {f3(b['p2'])}, pero Grad-CAM se calcula con su logit y aun así localiza al gato; el IoU = 0 también se ve favorecido porque los dos animales están separados por más de 3 celdas.")
c.para("**Límite de resolución.** Cada celda de $7\\times7$ equivale a $32\\times32$ píxeles de la entrada (224/7 = 32). Dos objetos que estén a menos de ≈ 32 px, o un objeto menor que una celda, caen en la misma celda o en celdas vecinas y no se pueden separar. "
       "Además, la interpolación bilineal solo suaviza el mapa y no crea información nueva: un objeto de 16×16 px sale como una mancha de al menos 32×32 y el recuadro que se infiere es siempre más grande que el objeto. "
       "Y el campo receptivo de layer4 abarca casi toda la imagen, por lo que el valor de una celda mezcla evidencia de objetos vecinos.")
code(c, "t1_4b")

# ============================================================ Task 1.4 (c)
cc = R["t1_4c"]; ts = cc["tests"]
c = C(39)
c.para(f"Clase explicada: {cc['cls']}. **Fracción de la masa del mapa normalizado dentro del recuadro: {f3(cc['mass_in'])}** ({100*cc['mass_in']:.1f}%), con un recuadro que ocupa solo el {100*cc['area_frac']:.1f}% de la imagen, "
       f"es decir {cc['ratio']:.1f} veces lo que tendría un mapa uniforme.")
c.para("**Prueba causal.** Se reemplazó por gris medio (0.5 en cada canal, antes de normalizar) todo lo que queda fuera del recuadro y, como control, lo que queda dentro:")
c.table(["Entrada", f"p({cc['cls']})", "Clase top-1 (prob.)"],
        [["Original", f4(ts["original"]["p_clase"]), f"{ts['original']['top1']} ({f3(ts['original']['p_top1'])})"],
         ["Fondo gris (solo queda el recuadro)", f4(ts["fondo gris (solo objeto)"]["p_clase"]), f"{ts['fondo gris (solo objeto)']['top1']} ({f3(ts['fondo gris (solo objeto)']['p_top1'])})"],
         ["Recuadro gris (solo queda el fondo)", f4(ts["objeto gris (solo fondo)"]["p_clase"]), f"{ts['objeto gris (solo fondo)']['top1']} ({f3(ts['objeto gris (solo fondo)']['p_top1'])})"]],
        "Tabla 7. Prueba causal en la imagen (c).")
c.image("t1_4c_prueba_causal.png", 6.6, "Figura 5. Grad-CAM con el recuadro del objeto y las dos pruebas causales.")
c.para(f"**¿Coinciden?** Solo en parte. El mapa dice que el modelo mira sobre todo al sofá ({100*cc['mass_in']:.0f}% de la masa en el 18% del área), pero un {100*(1-cc['mass_in']):.0f}% queda fuera del recuadro, en gran parte en el halo de pasto y sombra alrededor del sofá. "
       f"La prueba causal muestra que ambas partes son necesarias: sin el fondo, la probabilidad cae de {f3(ts['original']['p_clase'])} a {f3(ts['fondo gris (solo objeto)']['p_clase'])} (−{100*(1-ts['fondo gris (solo objeto)']['p_clase']/ts['original']['p_clase']):.0f}%), "
       f"y sin el objeto cae a {f4(ts['objeto gris (solo fondo)']['p_clase'])}. Es decir, el objeto es imprescindible y el contexto que lo rodea también contribuye, algo que el mapa por sí solo sugiere con su halo.")
c.para(f"**Conclusión sobre un atajo.** Hay evidencia de dependencia del contexto, porque el sofá aislado en su recuadro ya no basta para sostener la predicción, pero no de un atajo en el que el fondo decida la clase: con el sofá tapado, la clase bathtub desaparece (p ≈ {cc['tests']['objeto gris (solo fondo)']['p_clase']:.4f}). "
       f"La confusión parece venir de la forma del propio objeto (blanco, redondeado) y no del pasto; de hecho la clase correcta, studio couch, tiene p = {f4(cc['p_couch'])} (posición {cc['rank_couch']} de 1000). "
       "Hay que advertir que el parche gris introduce bordes rectangulares que el modelo no vio al entrenar, así que parte de la caída puede deberse a esa distorsión y no solo a quitar información: no se debe leer la prueba como una medida exacta de importancia.")
code(c, "t1_4c")

# ============================================================ Task 2.1
m = R["t2_1"]
c = C(48)
c.para(f"El conjunto (fetch_openml, credit-g versión 1) tiene {m['n']} solicitudes, {m['n_num']} variables numéricas y {m['n_cat']} categóricas, con 700 good (1) y 300 bad (0). "
       f"Las categóricas se codificaron con OrdinalEncoder, con el orden de categorías del propio conjunto. **Limitación:** esos códigos son arbitrarios (el 0, 1, 2… no refleja ningún orden de riesgo ni distancia entre categorías) y se tratan como numéricos. "
       "Un árbol puede aislar cualquier categoría con dos cortes, pero eso hace que (i) los colores del beeswarm de las categóricas (valor alto o bajo) no tengan significado, "
       "(ii) la comparación “valor mayor” carezca de sentido para ellas y (iii) la importancia de una categórica se reparta entre varios cortes.")
c.para(f"División estratificada 80/20 con semilla 42: {m['n_train']} instancias de entrenamiento y {m['n_test']} de prueba (con {100*m['good_train']:.0f}% y {100*m['good_test']:.0f}% de good, respectivamente). "
       "Modelo: GradientBoostingClassifier(n_estimators=200, max_depth=3, random_state=42), con el resto de parámetros por defecto (learning_rate = 0.1). "
       "Fondo para SHAP: 100 instancias de entrenamiento elegidas al azar (semilla 42).")
c.table(["Métrica (conjunto de prueba)", "Valor"], [["Exactitud", f4(m["acc"])], ["AUC", f4(m["auc"])]], "Tabla 8. Desempeño del modelo de crédito.")
code(c, "t2_1")

# ============================================================ Task 2.2
s = R["t2_2"]
c = C(52)
c.para("Se usó shap.TreeExplainer(modelo, data=fondo, feature_perturbation=\"interventional\", model_output=\"raw\"), cuya salida está en log-odds, sobre las 200 instancias de prueba. "
       f"**Eficiencia:** el error absoluto máximo de |valor base + Σφ_j − f(x)| sobre las 200 instancias es **{sci(s['eff_max'])}** (media {sci(s['eff_mean'])}), es decir, error de redondeo.")
code(c, "t2_2a")
c = C(54)
c.para(f"Valor base de SHAP = **{s['base']:.8f}**; promedio de la salida del modelo en log-odds sobre el fondo de 100 instancias = **{s['base_bg']:.8f}**. Son idénticos (diferencia {abs(s['base']-s['base_bg']):.1e}). "
       f"Equivale a σ(φ₀) = {R['t2_6']['p_base']:.3f}; no es la tasa de good del fondo, porque es la sigmoide del promedio de los log-odds y no el promedio de las probabilidades.")
c = C(56)
rk = lambda d, k: sorted(d, key=d.get, reverse=True)[:k]
ts_, tp_ = rk(s["I_shap"], 5), rk(s["I_perm"], 5)
c.para("Importancia global $I_j=\\frac1N\\sum_n|\\phi_j^{(n)}|$ (log-odds) e importancia por permutación (sklearn, scoring = roc_auc, 10 repeticiones, conjunto de prueba):")
c.table(["#", "Variable (|φ| medio)", "I_j", "Variable (permutación)", "Caída de AUC"],
        [[i + 1, ts_[i], f3(s["I_shap"][ts_[i]]), tp_[i], f3(s["I_perm"][tp_[i]])] for i in range(5)], "Tabla 9. Cinco variables principales de cada importancia.")
c.para(f"**Correlación de Spearman entre los dos rankings completos (20 variables): ρ = {s['spearman']:.3f}.** Coinciden en la variable principal (checking_status) y en que duration, credit_amount y savings_status están entre las seis primeras; "
       "difieren en credit_history (4.ª por permutación, 7.ª por SHAP) y en purpose y age, que SHAP considera importantes (5.ª y 6.ª) y la permutación casi no (valores ≈ 0.004 y 0.001).")
c = C(58)
ok, bd = s["top3_ok"], s["top3_bad"]
c.para(f"Se eligió, entre los clientes de prueba con p(good) > 0.8, el más cercano a 0.9, y entre los de p(good) < 0.2, el más cercano a 0.1. "
       f"**Aprobado:** instancia de prueba #{s['i_ok']}, p(good) = {f3(s['p_ok'])}. **Rechazado:** instancia #{s['i_bad']}, p(good) = {f3(s['p_bad'])}. Valor base $\\phi_0={s['base']:.3f}$ log-odds.")
c.image("t2_2_waterfall_aprobado.png", 5.6, "Figura 6. Cascada del cliente aprobado.")
c.para(f"Las tres mayores contribuciones del **aprobado** (en log-odds): (1) tener checking_status = {ok[0]['valor']} (sin cuenta corriente) aporta {ok[0]['phi']:+.2f}, la más grande; "
       f"(2) el propósito {ok[1]['valor']} (electrodoméstico) aporta {ok[1]['phi']:+.2f}; (3) una duración de {ok[2]['valor']} meses aporta {ok[2]['phi']:+.2f}. El resultado es f(x) = {s['fx_ok']:.2f} log-odds.")
c.image("t2_2_waterfall_rechazado.png", 5.6, "Figura 7. Cascada del cliente rechazado.")
c.para(f"Las tres mayores del **rechazado**: (1) una duración de {bd[0]['valor']} meses aporta {bd[0]['phi']:+.2f}; (2) un checking_status de {bd[1]['valor']} (saldo entre 0 y 200) aporta {bd[1]['phi']:+.2f}; "
       f"(3) un installment_commitment de {bd[2]['valor']} (la tasa de cuota más alta) aporta {bd[2]['phi']:+.2f}. Parten de $\\phi_0=1.29$ y terminan en f(x) = {s['fx_bad']:.2f} log-odds.")
c.image("t2_2_beeswarm.png", 5.8, "Figura 8. Beeswarm global de los valores SHAP de las 200 instancias de prueba.")
c.para("**Beeswarm.** checking_status es la variable con mayor dispersión (de ≈ −1 a +1.7), y sus dos grupos principales de puntos están muy separados. En credit_amount y duration los valores altos (rojo) empujan hacia log-odds negativos "
       "(créditos más grandes y largos son más riesgosos), con colas de hasta −2.6 y −1.5, y los bajos hacia positivos. Para las categóricas (checking_status, savings_status, purpose) el color solo muestra el código ordinal, "
       "que es arbitrario, por lo que no se debe leer como “más” o “menos”.")
code(c, "t2_2d")

# ============================================================ Task 2.3
r3 = R["t2_3"]; tb = r3["tab"]
c = C(61)
c.para(f"La variable numérica de mayor $I_j$ es **{r3['var']}** ($I_j={f3(tb['I_orig_sin_copia'])}$). Se creó {r3['var']}_copia $=x+0.05\\,\\sigma_x\\,\\epsilon$ con $\\sigma_x={r3['sigma']:.1f}$ (desviación estándar del entrenamiento) y $\\epsilon\\sim\\mathcal N(0,1)$ "
       f"(semilla 42), de modo que su correlación con la original es {r3['corr']:.4f}. El modelo de 21 variables, con la misma semilla e hiperparámetros, tiene exactitud {f4(r3['acc'])} y AUC {f4(r3['auc'])}. "
       "Resultados en el conjunto de prueba (log-odds):")
c.table(["Configuración", "Prom. |φ| original", "Prom. |φ| copia", "Prom. |φ_orig + φ_copia|"],
        [["Intervencional (fondo de 100)", f3(tb["interventional"]["orig"]), f3(tb["interventional"]["copia"]), f3(tb["interventional"]["suma"])],
         ["tree_path_dependent (sin fondo)", f3(tb["tree_path_dependent"]["orig"]), f3(tb["tree_path_dependent"]["copia"]), f3(tb["tree_path_dependent"]["suma"])],
         ["Referencia: I_j de la original en el modelo sin copia", f3(tb["I_orig_sin_copia"]), "—", "—"]], "Tabla 10. Reparto de crédito entre una variable y su copia con ruido.")
c.para(f"En las dos configuraciones el crédito se reparte entre la variable y su copia (≈ {f2(tb['interventional']['orig'])} y {f2(tb['interventional']['copia'])} en la intervencional). El efecto conjunto $|\\phi_{{orig}}+\\phi_{{copia}}|$ "
       f"({f3(tb['interventional']['suma'])} y {f3(tb['tree_path_dependent']['suma'])}) queda cerca, algo por debajo, de la importancia de la variable sola ({f3(tb['I_orig_sin_copia'])}).")
code(c, "t2_3")

# ============================================================ Task 2.4
r4 = R["t2_4"]["rows"]
c = C(65)
c.para(f"KernelSHAP (shap.KernelExplainer sobre la salida en log-odds de decision_function, nsamples = {R['t2_4']['n_samples']}, silent) sobre las primeras {R['t2_4']['n_inst']} instancias de prueba, con fondos de 10, 50 y 100 instancias "
       "(los primeros 10, 50 y 100 del fondo del Task 2.1). Los valores resultantes se compararon con TreeSHAP en el apartado siguiente.")
code(c, "t2_4")
c = C(67)
c.table(["Fondo", "Error abs. medio vs TreeSHAP", "Error relativo (÷ prom. |φ|)", "Tiempo KernelSHAP (10 inst.)", "Tiempo TreeSHAP (10 inst.)"],
        [[x["fondo"], f4(x["mae"]), f"{100*x['mae_rel']:.1f}%", f"{x['t_kernel']:.2f} s", f"{1000*x['t_tree']:.1f} ms"] for x in r4], "Tabla 11. KernelSHAP frente a TreeSHAP intervencional con el mismo fondo.")
c.para(f"El error medio (≈ 0.04 log-odds, ≈ 16% del |φ| medio) no cambia de forma apreciable con el fondo. El tiempo de KernelSHAP crece con el fondo ({r4[0]['t_kernel']:.2f} → {r4[1]['t_kernel']:.2f} → {r4[2]['t_kernel']:.2f} s) y es entre {min(x['t_kernel']/x['t_tree'] for x in r4):.0f} y {max(x['t_kernel']/x['t_tree'] for x in r4):.0f} veces mayor que el de TreeSHAP.")

# ============================================================ Task 2.5
r5 = R["t2_5"]; F6 = r5["feats"]
c = C(71)
c.para("Se entrenó un segundo modelo (200 árboles, profundidad 3, semilla 42) solo con duration, credit_amount, installment_commitment, residence_since, age y existing_credits, con un fondo de 50 instancias de entrenamiento aleatorias "
       "y las primeras 5 instancias de prueba. La función shapley_exact del notebook construye para cada instancia la matriz con las $2^6=64$ coaliciones y evalúa el modelo una sola vez por coalición y fondo. "
       "$v(S)$ es el promedio sobre el fondo de $f(x_S,b_{\\bar S})$, donde las variables de $S$ se fijan en el valor de la instancia, las demás se toman del fondo y $f$ es la salida en log-odds.")
code(c, "t2_5")
c = C(73)
c.para("Para cada variable $i$ se sumó, sobre las 32 coaliciones $S\\subseteq N\\setminus\\{i\\}$, el aporte marginal con el peso de factoriales:")
c.eq(r"\phi_i=\sum_{S\subseteq N\setminus\{i\}}\frac{|S|!\,(n-|S|-1)!}{n!}\,\big[v(S\cup\{i\})-v(S)\big],\qquad n=6")
c.table(["Inst."] + F6 + ["v(N) = f(x)", "v(∅)"],
        [[x["inst"]] + [f"{x[f]:+.4f}" for f in F6] + [f"{x['v_N']:.4f}", f"{x['v_vacio']:.4f}"] for x in r5["rows"]], "Tabla 12. Valores de Shapley exactos (log-odds) calculados a mano.")
c = C(75)
c.para(f"Comparados con shap.TreeExplainer intervencional con el mismo fondo de 50: **error absoluto máximo = {sci(r5['err_max'])}** (sobre las 5 instancias y 6 variables), es decir, coinciden hasta el redondeo de float32. "
       f"**Eficiencia:** $\\left|\\sum_i\\phi_i-(v(N)-v(\\emptyset))\\right|$ tiene máximo {sci(r5['eff_max'])}, y $v(N)=f(x)$ en todas las instancias (Tabla 12). "
       "TreeSHAP intervencional calcula exactamente la misma cantidad que la fórmula de Shapley con $v(S)$ definido sobre el fondo, pero en tiempo polinomial.")

# ============================================================ Task 2.6 (a)
r6 = R["t2_6"]
c = C(79)
ep = sorted(r6["eff_p"].items(), key=lambda kv: kv[1])
c.para(f"Cliente rechazado (#{s['i_bad']}): $f(x)={r6['f_bad']:.3f}$ log-odds y $\\phi_0={r6['base']:.3f}$.")
c.table(["Cantidad", "Valor"],
        [["p(x) = σ(f(x))", f4(r6["p_x"])], ["Probabilidad base σ(φ₀)", f4(r6["p_base"])], ["p(x) − σ(φ₀)", f4(r6["p_x"] - r6["p_base"])],
         ["Σ_j [σ(φ₀ + φ_j) − σ(φ₀)]", f4(r6["sum_eff_p"])]], "Tabla 13. Efecto en probabilidad del cliente rechazado.")
c.para("Mayores efectos individuales en probabilidad: " + ", ".join(f"{k} {v:+.3f}" for k, v in ep[:4]) + "; los positivos más grandes son " + ", ".join(f"{k} {v:+.3f}" for k, v in ep[::-1][:2]) + ".")
c.para(f"La suma de efectos ({f4(r6['sum_eff_p'])}) **no** es igual a $p(x)-\\sigma(\\phi_0)$ ({f4(r6['p_x']-r6['p_base'])}); difieren en {abs(r6['sum_eff_p']-(r6['p_x']-r6['p_base'])):.3f}. "
       "**Por qué.** La propiedad de eficiencia de Shapley se cumple sobre la función que se explica, que aquí es $f$ (log-odds): $\\phi_0+\\sum_j\\phi_j=f(x)$ exactamente. La probabilidad es $\\sigma(f)$ y $\\sigma$ es no lineal, de modo que")
c.eq(r"\sum_j\left[\sigma(\phi_0+\phi_j)-\sigma(\phi_0)\right]\neq\sigma\Big(\phi_0+\sum_j\phi_j\Big)-\sigma(\phi_0)")
c.para("Cada término del lado izquierdo mide el efecto de la variable $j$ sola, partiendo del valor base, ignorando que las demás también actúan, mientras que el lado derecho aplica $\\sigma$ una vez a la suma total. "
       "Aquí cada efecto individual se mide partiendo de $\\phi_0=1.29$, donde $\\sigma$ es empinada, mientras que el efecto conjunto termina en la cola plana de $\\sigma$ (p = 0.10); por eso la suma de efectos individuales sobrestima la caída total. "
       "Para sumar exacto en probabilidad habría que aplicar SHAP a $\\sigma\\circ f$ como función.")

# ============================================================ Task 2.6 (b)
c = C(81)
vr = r6["v_rank"]
c.para(f"La mayor diferencia de posición es **{vr}**: ocupa el puesto {r6['rk_s']} de 20 por $I_j$ ($|\\phi|$ medio = {f3(r6['shap_v'])}) y el puesto {r6['rk_p']} por permutación "
       f"(caída de AUC = {r6['perm_te']:+.4f} ± {r6['perm_te_std']:.4f}, es decir, estadísticamente cero). Las siguientes mayores diferencias son purpose y age (7 puestos cada una).")
c.para("**Hipótesis 1: la variable está correlacionada con otras y la permutación se “compensa”.** No está respaldada: sus correlaciones de Spearman con las demás variables son bajas, " +
       ", ".join(f"{k} = {v:+.3f}" for k, v in r6["corr_top"].items()) + " (todas con |ρ| < 0.24).")
c.para(f"**Hipótesis 2 (respaldada): el modelo usa la variable, pero no aporta señal que generalice.** Su importancia por permutación en **entrenamiento** es {r6['perm_train']:.4f} ± {r6['perm_train_std']:.4f} (positiva y significativa), "
       f"y la de impureza (feature_importances_) es {r6['gini']:.3f}: el modelo sí hizo cortes sobre ella al entrenar. Pero la variable no predice la clase: su AUC univariado es {r6['auc_uni_tr']:.3f} en entrenamiento y {r6['auc_uni_te']:.3f} en prueba, "
       f"y su correlación de Spearman con la clase en prueba es {r6['spearman_y']:+.3f}. Es decir, los 200 árboles ajustaron ruido en entrenamiento (sobreajuste), y en prueba barajarla no cambia el AUC.")
c.para("**Qué mide cada importancia.** La importancia por permutación mide cuánto **cae una métrica de desempeño** (AUC) al romper la relación entre la variable y la salida real: es una medida de la utilidad predictiva de la variable "
       "para el modelo, sobre datos de prueba. $I_j=\\frac1N\\sum|\\phi_j|$ mide el **tamaño promedio de su contribución a la salida** del modelo en log-odds, sin mirar si esa contribución acierta. "
       "Por eso una variable puede mover mucho la salida (SHAP alto) y no mejorar el AUC (permutación ≈ 0): el modelo la usa, pero en esos usos se compensa o no ayuda a acertar.")

# ============================================================ Task 2.6 (c)
c = C(83)
tI, tT = tb["interventional"], tb["tree_path_dependent"]
c.para(f"**Qué pasa con el crédito.** Sin la copia, {r3['var']} tiene $I_j={f3(tb['I_orig_sin_copia'])}$. Con la copia, el crédito se divide entre las dos: {f3(tI['orig'])} y {f3(tI['copia'])} en la intervencional, {f3(tT['orig'])} y {f3(tT['copia'])} en la dependiente de la trayectoria. "
       f"La suma de valores absolutos ({f3(tI['orig']+tI['copia'])}) supera al valor del par $|\\phi_{{orig}}+\\phi_{{copia}}|$ ({f3(tI['suma'])}) porque en el {100*tI['frac_signo_opuesto']:.0f}% de las instancias los dos valores tienen signo opuesto ({100*tT['frac_signo_opuesto']:.0f}% en la dependiente de la trayectoria): "
       "los árboles aprovechan la pequeña diferencia entre las dos variables. Cada una de las dos queda con menos crédito que la original sola, y el par junto queda cerca (algo por debajo) del valor de la variable original: el crédito del “concepto” se conserva casi, pero se fragmenta entre sus proxies.")
c.para("**Axioma.** El de **simetría** (variables que contribuyen igual a toda coalición deben recibir el mismo crédito). No reparte por igual porque la premisa no se cumple: la copia no es idéntica (lleva ruido de 5% de $\\sigma_x$) y, "
       "sobre todo, el modelo no es simétrico en las dos variables. Los árboles son voraces y en cada nodo eligen la que da el mejor corte, que varía de un árbol a otro. Ni la variable ni su copia son intercambiables para $f$, "
       "así que Shapley no obliga a que valgan lo mismo. Además, en la versión intervencional $v(S)$ combina el valor de una variable de la instancia con el de la otra tomado del fondo (una combinación casi imposible con correlación 0.9988), con lo que se evalúa el modelo en regiones donde no tiene datos y el reparto depende de cómo extrapolan los árboles.")
c.para(f"**Comparación de las dos configuraciones y por qué tree_path_dependent no es la condicional.** Ambas dan resultados parecidos. La dependiente de la trayectoria asigna algo menos a ambas ({f3(tT['orig'])} y {f3(tT['copia'])} contra {f3(tI['orig'])} y {f3(tI['copia'])}) y reparte menos por igual (copia/original = {tT['copia']/tT['orig']:.2f} contra {tI['copia']/tI['orig']:.2f}). "
       "Esta versión no usa fondo: cuando una variable falta de la coalición, el algoritmo manda la muestra por las dos ramas del árbol según la proporción de datos de entrenamiento (cover) de cada rama. "
       "Eso aproxima la esperanza condicional, pero solo con la información de la estructura de cada árbol. No condiciona sobre el valor de las demás variables de la coalición: si el árbol corta por la original, no usa que la copia (conocida) permite predecir el valor de la original. "
       "La condicional exacta $E[f(X)\\mid X_S=x_S]$ sí lo haría y vería a las dos variables como casi la misma información, de modo que el resultado no es igual a ella.")
c.para("**Implicación en el mundo real.** Un valor SHAP describe cómo **este modelo** usa cada variable y no la importancia de la variable en el mundo. Con variables correlacionadas el crédito se reparte de forma arbitraria "
       "(depende de la semilla, de los cortes que eligieron los árboles y de la versión de SHAP), así que una variable realmente importante puede parecer poco importante si tiene proxies, y una redundante puede parecer importante. "
       "Antes de interpretarlo conviene revisar correlaciones, agrupar variables redundantes y no leer $\\phi$ como un efecto causal (“si el cliente redujera el monto, la probabilidad subiría”).")

# ============================================================ Task 2.6 (d)
c = C(85)
ev = r6["evals"]; k = R["t2_4"]["rows"]
ratio_t = [k[i]["t_kernel"] / k[0]["t_kernel"] for i in range(3)]
c.para("**Modelo de costo de KernelSHAP.** Con nsamples muestras y un fondo de $B$ instancias se evalúa el modelo $\\text{nsamples}\\times B$ veces:")
c.table(["Fondo B", "Evaluaciones (2 000 × B)", "Tiempo (10 inst.)", "Tiempo relativo a B = 10", "Evaluaciones relativas a B = 10", "Tiempo por 10⁶ evaluaciones"],
        [[k[i]["fondo"], f"{ev[i]:,}", f"{k[i]['t_kernel']:.2f} s", f"{ratio_t[i]:.1f}×", f"{ev[i]/ev[0]:.0f}×", f"{1e6*k[i]['t_kernel']/(ev[i]*R['t2_4']['n_inst']):.2f} s"] for i in range(3)],
        "Tabla 14. Costo de KernelSHAP.")
c.para(f"**¿Crece el tiempo como el modelo?** Aproximadamente sí. Al pasar de B = 10 a 100 las evaluaciones se multiplican por 10 y el tiempo por {ratio_t[2]:.1f}. El tiempo no es proporcional del todo porque hay costos fijos "
       "(construcción del explicador, resolución de la regresión ponderada, llamadas a predict) que pesan más con fondos chicos; con fondos grandes el costo de las evaluaciones domina y el tiempo por evaluación se estabiliza. "
       "**¿Baja el error al aumentar el fondo?** No: " + ", ".join(f"{x['mae']:.4f}" for x in k) + " para B = 10, 50 y 100. La referencia TreeSHAP se calcula con el mismo fondo, así que el fondo no introduce diferencia. "
       "El error que queda es de muestreo de las 2 000 coaliciones de entre $2^{20}$, que depende del número de muestras y no del fondo.")
c.para(f"**Shapley exacto con 20 variables.** Una predicción sobre 100 instancias (el costo de un $v(S)$ con el fondo de 100) tarda {1000*r6['t100']:.3f} ms (promedio de 200 repeticiones). Las $2^{{20}}=1\\,048\\,576$ evaluaciones de $v(S)$ para una sola instancia tomarían "
       f"$2^{{20}}\\times{1000*r6['t100']:.3f}$ ms = **{r6['t_exact_20']:.0f} s ≈ {r6['t_exact_20']/60:.1f} minutos**. KernelSHAP tarda {r6['t_kernel_inst']:.2f} s por instancia con el fondo de 100, "
       f"unas {r6['t_exact_20']/r6['t_kernel_inst']:,.0f} veces menos, a cambio de una aproximación (error ≈ 0.04 log-odds). TreeSHAP lo calcula exactamente en {1000*k[2]['t_tree']/10:.1f} ms por instancia.")
c.para(f"**Con 60 variables:** $2^{{60}}\\approx1.15\\times10^{{18}}$ evaluaciones $\\times\\,10^{{-6}}$ s = $1.15\\times10^{{12}}$ s ≈ **{r6['years_60']:,.0f} años** ($3.65\\times10^{{4}}$). Es inviable, por lo que para modelos con muchas variables hay que usar métodos específicos (TreeSHAP) o aproximaciones por muestreo (KernelSHAP).")
code(c, "t2_6")

doc.save(DST)
print("guardado:", DST)
